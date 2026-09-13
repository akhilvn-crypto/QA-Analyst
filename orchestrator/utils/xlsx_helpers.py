"""Shared low-level XLSX formatting helpers -- the openpyxl analogue of
docx_helpers.py, for any generation script that writes a controlled-document
workbook (Document Version Control / Document Release History front matter,
plus a tabular working sheet).

Extracted from zephyr_export.py (the original, single-caller implementation)
so any later caller needing the same cell-styling/row-height/logo-centering
logic can reuse it rather than duplicating ~300 lines of formatting code --
exactly the kind of drift risk .claude/rules/orchestrator-structure.md's
utils/ rule exists to prevent. zephyr_export.py re-imports these under its
original private names so its existing behaviour and test suite are
unchanged.
"""

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.worksheet import Worksheet

from orchestrator.utils.paths import project_logo_path

HEADER_FILL = "1F3864"
HEADER_FONT_COLOR = "FFFFFF"
DEFAULT_COLUMN_WIDTH = 16
POINTS_PER_LINE = 15
LOGO_WIDTH_PX = 150

_THIN_SIDE = Side(style="thin", color="000000")
THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)

# Real blank margin -- a genuine gap of unset-width columns/rows before an
# audit table's header, so the table itself doesn't start at the sheet's A1
# corner. This is the actual centering mechanism for the two front-matter
# sheets: physically moving where the table starts, not a print/view setting
# that only changes how the file looks when printed.
AUDIT_LEFT_MARGIN_COLS = 3
AUDIT_TOP_MARGIN_ROWS = 3


def estimated_line_count(text: str, column_width: float) -> int:
    """How many visual lines `text` will wrap onto in a column of
    `column_width` (openpyxl/Excel character-width units), so a row can be
    sized tall enough to show all of it -- Excel wraps a single long line
    across several visual rows within one cell, and a height computed from
    raw '\\n' counts alone under-sizes any row containing one of those,
    silently clipping it."""
    if not text:
        return 1
    # A small safety margin below the raw width: cell padding and the
    # occasional wider character mean the true wrap point is a little
    # narrower than the nominal column width -- underestimating chars-per-
    # line (and so overestimating lines needed) fails safe by making the
    # row slightly taller than strictly necessary, never clipped.
    chars_per_line = max(1, int(column_width) - 2)
    total = 0
    for line in text.split("\n"):
        total += max(1, -(-len(line) // chars_per_line))  # ceil division
    return max(1, total)


def row_height(
    row: list[str], columns: list[str], wrap_columns: set[str], column_widths: dict[str, float]
) -> float:
    """The row height needed to show every wrap-enabled cell in `row`
    without clipping -- the max estimated line count across ALL of them,
    not just one column (a long free-text sentence with no embedded newline
    still wraps and needs the same consideration as a multi-line cell)."""
    max_lines = 1
    for idx, column in enumerate(columns):
        if column not in wrap_columns or idx >= len(row):
            continue
        width = column_widths.get(column, DEFAULT_COLUMN_WIDTH)
        max_lines = max(max_lines, estimated_line_count(row[idx], width))
    return max(POINTS_PER_LINE, POINTS_PER_LINE * max_lines)


def populate_sheet(
    sheet: Worksheet,
    columns: list[str],
    rows: list[list[str]],
    *,
    wrap_columns: set[str],
    column_widths: dict[str, float],
    add_borders: bool = False,
    center_columns: set[str] = frozenset(),
    start_row: int = 1,
    start_col: int = 1,
) -> None:
    """Shared formatting for a tabular XLSX sheet: bold navy header, frozen
    header row, per-column width, per-row height sized to its tallest
    wrapped cell, wrap-text alignment on the text columns, and a filter
    dropdown across the header -- this is what keeps every row fully
    visible regardless of how much text it holds, instead of some rows
    happening to fit and others rendering clipped.

    `add_borders`/`center_columns` are opt-in (default off) -- used by
    audit-facing front-matter sheets, which need a bordered, centered-text
    table, not by an ordinary working sheet.

    `start_row`/`start_col` place the table's header at that cell instead
    of always A1 -- this is the actual "center the table" mechanism: a
    genuine blank margin of real rows/columns before the table, not a
    print/view setting that only affects how the file looks when printed."""
    header_row = start_row
    header_fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
    for idx, column in enumerate(columns):
        col_idx = start_col + idx
        cell = sheet.cell(row=header_row, column=col_idx, value=column)
        cell.font = Font(bold=True, color=HEADER_FONT_COLOR)
        cell.fill = header_fill
        cell.alignment = (
            Alignment(horizontal="center", vertical="center", wrap_text=True)
            if add_borders
            else Alignment(vertical="center", wrap_text=True)
        )
        sheet.column_dimensions[get_column_letter(col_idx)].width = column_widths.get(
            column, DEFAULT_COLUMN_WIDTH
        )

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=start_col).coordinate

    last_row = header_row
    for row in rows:
        last_row += 1
        for idx, value in enumerate(row):
            column = columns[idx]
            cell = sheet.cell(row=last_row, column=start_col + idx, value=value)
            horizontal = "center" if column in center_columns else None
            cell.alignment = Alignment(
                horizontal=horizontal, wrap_text=column in wrap_columns, vertical="top"
            )
        sheet.row_dimensions[last_row].height = row_height(row, columns, wrap_columns, column_widths)

    first_col_letter = get_column_letter(start_col)
    last_col_letter = get_column_letter(start_col + len(columns) - 1)
    sheet.auto_filter.ref = f"{first_col_letter}{header_row}:{last_col_letter}{last_row}"

    if add_borders:
        apply_borders(
            sheet,
            min_row=header_row,
            max_row=last_row,
            min_col=start_col,
            max_col=start_col + len(columns) - 1,
        )


def apply_borders(sheet: Worksheet, *, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    """Thin black border on all four sides of every cell in the given
    range -- an unbordered range of cell writes shows no grid at all
    on-screen or when printed, unlike a real Excel Table or a Word
    'Table Grid' style."""
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = THIN_BORDER


def column_width_to_px(width: float) -> int:
    """Approximate pixel width of an Excel column given its openpyxl
    character-unit `width` -- the standard ECMA-376 approximation for the
    default font (width_px ~= chars*7 + 5), close enough for centering
    (which is inherently approximate -- exact rendering varies by font/DPI
    anyway) without depending on a fixed-size assumption."""
    return round(width * 7) + 5


def centered_image_anchor(
    span_columns: list[str],
    column_widths: dict[str, float],
    *,
    row_idx: int,
    width_px: int,
    height_px: int,
) -> OneCellAnchor:
    """A `OneCellAnchor` positioning an image horizontally centered across
    `span_columns` -- an image is a floating object, not a cell, so
    centering it requires computing where its centered left edge actually
    falls, then expressing that as a (column, pixel-offset-within-column)
    pair. Anchoring at a fixed cell like "B2" would always left-align it to
    that cell's edge instead."""
    span_widths_px = [column_width_to_px(column_widths.get(c, DEFAULT_COLUMN_WIDTH)) for c in span_columns]
    total_px = sum(span_widths_px)
    offset_px = max(0, (total_px - width_px) // 2)

    col_idx = column_index_from_string(span_columns[0]) - 1  # 0-based
    remaining = offset_px
    for col_width_px in span_widths_px:
        if remaining < col_width_px:
            break
        remaining -= col_width_px
        col_idx += 1

    anchor = OneCellAnchor()
    anchor._from = AnchorMarker(col=col_idx, colOff=pixels_to_EMU(remaining), row=row_idx, rowOff=0)
    anchor.ext = XDRPositiveSize2D(cx=pixels_to_EMU(width_px), cy=pixels_to_EMU(height_px))
    return anchor


def add_logo(
    sheet: Worksheet,
    *,
    span_columns: list[str],
    column_widths: dict[str, float],
    row_idx: int,
    width_px: int = LOGO_WIDTH_PX,
) -> None:
    """Embeds the project logo horizontally centered across `span_columns`,
    scaled to `width_px` wide with its original aspect ratio preserved. A
    missing logo asset is a no-op, not a failure -- the same tolerance
    report_exporter.py and test_plan_docx_writer.py give a missing logo
    file."""
    logo_path = project_logo_path()
    if not logo_path.is_file():
        return
    image = XLImage(str(logo_path))
    aspect_ratio = image.height / image.width
    height_px = round(width_px * aspect_ratio)
    image.anchor = centered_image_anchor(
        span_columns, column_widths, row_idx=row_idx, width_px=width_px, height_px=height_px
    )
    sheet.add_image(image)


def write_version_control_sheet(
    workbook: Workbook,
    document_control,
    *,
    fields: list[tuple[str, str]],
    center_fields: set[str] = frozenset(),
    label_width: float = 22,
    value_width: float = 55,
    sheet_name: str = "Document Version Control",
) -> Worksheet:
    """Writes a "Document Version Control" sheet: centered project logo,
    heading, then a bordered label/value form for every (label, attr) pair
    in `fields`, reading `attr` off `document_control` (any object with
    those attributes -- DocumentControlMeta in requirement.py/test_case.py/
    test_plan.py are all duck-type compatible). `center_fields` names which
    labels' values render centered (short categorical values); anything
    else stays left-aligned (free text like Description)."""
    sheet = workbook.create_sheet(sheet_name)
    label_col = AUDIT_LEFT_MARGIN_COLS + 1
    value_col = label_col + 1
    label_letter = get_column_letter(label_col)
    value_letter = get_column_letter(value_col)
    sheet.column_dimensions[label_letter].width = label_width
    sheet.column_dimensions[value_letter].width = value_width
    column_widths = {label_letter: label_width, value_letter: value_width}

    logo_row_idx = AUDIT_TOP_MARGIN_ROWS  # 0-based
    add_logo(
        sheet,
        span_columns=[label_letter, value_letter],
        column_widths=column_widths,
        row_idx=logo_row_idx,
    )

    row = AUDIT_TOP_MARGIN_ROWS + 9  # leaves room for the logo above
    heading = sheet.cell(row=row, column=label_col, value=sheet_name)
    heading.font = Font(bold=True, size=14, color=HEADER_FILL)
    heading.alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.merge_cells(start_row=row, start_column=label_col, end_row=row, end_column=value_col)
    row += 2

    first_field_row = row
    for label, attr in fields:
        label_cell = sheet.cell(row=row, column=label_col, value=label)
        label_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

        value_cell = sheet.cell(row=row, column=value_col, value=getattr(document_control, attr))
        centered = label in center_fields
        value_cell.alignment = Alignment(
            horizontal="center" if centered else None, wrap_text=True, vertical="top"
        )
        row += 1
    last_field_row = row - 1

    apply_borders(sheet, min_row=first_field_row, max_row=last_field_row, min_col=label_col, max_col=value_col)
    return sheet


def build_release_history_rows(release_history) -> list[list[str]]:
    """Row-ify a list of ReleaseHistoryEntry-shaped objects (any object with
    version/date/author/reviewed_by/reviewed_on/approved_by/approved_on/
    reasons attributes -- the test_case.py/test_plan.py dataclasses are all
    duck-type compatible) into the standard eight-column Document Release
    History row shape."""
    return [
        [
            entry.version,
            entry.date,
            entry.author,
            entry.reviewed_by,
            entry.reviewed_on,
            entry.approved_by,
            entry.approved_on,
            entry.reasons,
        ]
        for entry in release_history
    ]


def write_release_history_sheet(
    workbook: Workbook,
    rows: list[list[str]],
    *,
    columns: list[str],
    wrap_columns: set[str],
    column_widths: dict[str, float],
    center_columns: set[str] = frozenset(),
    sheet_name: str = "Document Release History",
) -> Worksheet:
    """Writes a "Document Release History" sheet -- a bordered, centered
    tabular sheet offset by the same AUDIT_LEFT_MARGIN_COLS/
    AUDIT_TOP_MARGIN_ROWS margin as the Version Control sheet, so the two
    front-matter sheets read as one consistent unit."""
    sheet = workbook.create_sheet(sheet_name)
    populate_sheet(
        sheet,
        columns,
        rows,
        wrap_columns=wrap_columns,
        column_widths=column_widths,
        add_borders=True,
        center_columns=center_columns,
        start_row=AUDIT_TOP_MARGIN_ROWS + 1,
        start_col=AUDIT_LEFT_MARGIN_COLS + 1,
    )
    return sheet
