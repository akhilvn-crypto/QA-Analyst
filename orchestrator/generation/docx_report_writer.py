"""Writes structured requirement-analysis data as a formatted Word document.

Usage:
    python -m orchestrator.generation.docx_report_writer <doc-name>

Reads output/requirement-analysis/<doc-name>-analysis.json and writes
output/requirement-analysis/<doc-name>-analysis.docx

The report is organized as five numbered sections -- 1. Document Control &
Metadata (plus its 1.1 Revision History subsection), 2. Project Overview &
Scope, 3. Functional Requirements Analysis & Acceptance Criteria, 4.
Non-Functional Requirements (NFR) Analysis, 5. Compliance & Regulatory
Requirements Analysis -- with sections 3-5 grouping the analyzed
requirements by their `category` field and rendering each as its own
full-width "card" (a heading, then labeled fields as normal paragraphs)
rather than a row in a wide table -- squeezing seven columns' worth of free
text into ~1.5-2in-wide table cells on a landscape page was reported as hard
to read, and this document is portrait now that column width no longer
needs reclaiming. See the output-structure skill for the full layout
rationale, including why sections 3-5 exist as separate categories instead
of one flat requirement list with two document-level narrative paragraphs
(the pre-`category` shape).
"""

import re
import sys
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from orchestrator.models.requirement import AnalysisDocument, DocumentControlMeta, ReleaseHistoryEntry, Requirement
from orchestrator.utils.docx_helpers import (
    FONT_NAME,
    add_bottom_border,
    add_header_table,
    add_key_value_table,
    add_page_field,
    clear_paragraph_border,
    save_document,
    set_authorship,
)
from orchestrator.utils.io_helpers import read_json
from orchestrator.utils.paths import analysis_docx_path, analysis_json_path, header_logo_path, project_logo_path
from orchestrator.validation.analysis_validator import (
    CATEGORY_COMPLIANCE,
    CATEGORY_FUNCTIONAL,
    CATEGORY_NON_FUNCTIONAL,
    STATUS_BLOCKED,
)
from orchestrator.validation.validate import validate_file

FONT_SIZE_PT = 12

ACCENT_COLOR = "1F3864"  # dark navy
HEADER_FILL = ACCENT_COLOR
HEADER_TEXT_COLOR = "FFFFFF"

# Status, colored the same way a reviewer would highlight it by hand --
# green/amber/red/grey, the standard traffic-light reading order from "no
# client input needed" to "blocked", plus a neutral grey for a Compliance
# item the inferred domain doesn't implicate at all. Keyed by the *display*
# status (see `_display_status`), not the raw JSON value -- STATUS_BLOCKED's
# raw value is the longer "Not Generated – Client Clarification Required".
STATUS_COLORS = {
    "Generated": "006100",
    "Generated with Assumptions": "9C6500",
    "Blocked": "C00000",
    "Not Applicable": "6B6B6B",
}

# (label, DocumentControlMeta attribute) pairs, in display order, for the
# "1. Document Control & Metadata" section. "Project Name"/"Date" are the
# report-facing labels for the `title`/`prepared_date` JSON fields.
DOCUMENT_CONTROL_FIELDS = [
    ("Project Name", "title"),
    ("Project ID", "project_id"),
    ("Document ID", "document_id"),
    ("Description", "description"),
    ("Document Version", "version"),
    ("Prepared By", "prepared_by"),
    ("Date", "prepared_date"),
    ("Approved Date", "approved_date"),
    ("Master Template ID", "master_template_id"),
    ("Classification", "classification"),
]

# "1.1 Revision History" is a condensed view of `release_history` -- Version/
# Date/Description(=reasons)/Author/Reviewed By/Approved By -- dropping
# Reviewed On/Approved On from this table only; the full eight-field entry
# is still what's stored in the JSON (see ReleaseHistoryEntry), unchanged.
REVISION_HISTORY_COLUMNS = ["Version", "Date", "Description", "Author", "Reviewed By", "Approved By"]
# Inches, summing to the portrait page's ~6.5in usable width (8.5in - 2*1in
# margins). Description gets by far the largest share, same reasoning as the
# eight-column table this condenses.
REVISION_HISTORY_WIDTHS_IN = [0.65, 0.85, 2.55, 0.95, 0.75, 0.75]
REVISION_HISTORY_FONT_SIZE_PT = 9.5

NO_GAPS_TEXT = "No significant gaps identified."
NO_QUESTION_TEXT = "None."
NO_RECOMMENDATIONS_TEXT = "No recommendations."
NO_REQUIREMENTS_TEXT = "No requirements were extracted from this document."

SECTION_TITLES = {
    CATEGORY_FUNCTIONAL: "3. Functional Requirements Analysis & Acceptance Criteria",
    CATEGORY_NON_FUNCTIONAL: "4. Non-Functional Requirements (NFR) Analysis",
    CATEGORY_COMPLIANCE: "5. Compliance & Regulatory Requirements Analysis",
}
NO_CATEGORY_REQUIREMENTS_TEXT = {
    CATEGORY_FUNCTIONAL: "No functional requirements were extracted from this document.",
    CATEGORY_NON_FUNCTIONAL: "No non-functional requirements were identified in this document.",
    CATEGORY_COMPLIANCE: "No compliance or regulatory considerations were identified as applicable to this document.",
}


def _display_status(status: str) -> str:
    """The report shows a shorter, bullet-friendly label than the JSON's own
    canonical `acceptance_criteria_status` value -- e.g. "Blocked" instead of
    "Not Generated – Client Clarification Required". The JSON value itself
    is unchanged (other agents/validation still key off the full string);
    this mapping is display-only. "Not Applicable" (Compliance-only) already
    matches its own display form, so it passes through unchanged."""
    return "Blocked" if status == STATUS_BLOCKED else status


def _format_gap_text(requirement: Requirement) -> str:
    if not requirement.gaps:
        return NO_GAPS_TEXT
    if len(requirement.gaps) == 1:
        return requirement.gaps[0].description
    return "\n".join(
        f"{idx}. {gap.description}" for idx, gap in enumerate(requirement.gaps, start=1)
    )


def _format_question_text(requirement: Requirement) -> str:
    if not requirement.gaps:
        return NO_QUESTION_TEXT
    if len(requirement.gaps) == 1:
        return requirement.gaps[0].question
    return "\n".join(
        f"{idx}. {gap.question}" for idx, gap in enumerate(requirement.gaps, start=1)
    )


def _format_recommendations_text(requirement: Requirement) -> str:
    """One line per recommendation (not a multi-line Recommendation/Reason/
    Business Benefit/Client Confirmation Recommended block) -- the common
    case of zero or one recommendation then renders as a single bullet line
    in both reports, matching how every other single-valued field does."""
    if not requirement.recommendations:
        return NO_RECOMMENDATIONS_TEXT
    lines = []
    for rec in requirement.recommendations:
        confirmation = "Yes" if rec.client_confirmation_recommended else "No"
        lines.append(
            f"{rec.recommendation} (Reason: {rec.reason}; "
            f"Business Benefit: {rec.business_benefit}; "
            f"Client Confirmation Recommended: {confirmation})"
        )
    return "\n".join(lines)


def _req_sort_key(req_id: str) -> tuple[str, int]:
    """Sort REQ-2 before REQ-10 (numeric, not lexicographic) regardless of
    zero-padding width, while still sorting non-numeric IDs predictably."""
    match = re.search(r"(\d+)$", req_id)
    if not match:
        return (req_id, 0)
    return (req_id[: match.start()], int(match.group(1)))


def _build_summary_line(doc_name: str, requirements: list[Requirement]) -> str:
    """One condensed line of report-wide statistics, shown at the end of the
    "1. Document Control & Metadata" section rather than as its own numbered
    section (the five-section layout this report follows has no room for a
    separate "Analysis Summary" heading)."""
    total = len(requirements)
    total_gaps = sum(len(req.gaps) for req in requirements)

    status_counts: dict[str, int] = {}
    for req in requirements:
        label = _display_status(req.acceptance_criteria_status)
        status_counts[label] = status_counts.get(label, 0) + 1
    status_line = " | ".join(f"{status}: {count}" for status, count in status_counts.items()) or "(none)"

    return (
        f"Document: {doc_name} | Generated: {date.today().isoformat()} | "
        f"Total requirements: {total} | Total gaps identified: {total_gaps} | "
        f"Acceptance Criteria: {status_line}"
    )


def _set_portrait_letter(document: Document) -> None:
    section = document.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)


def _set_base_styles(document: Document) -> None:
    heading_sizes = {"Title": 20, "Heading 1": 14, "Heading 2": 13}
    for style_name, size_pt in heading_sizes.items():
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size_pt)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(ACCENT_COLOR)
    # python-docx's built-in Title style ships with a bottom paragraph
    # border -- strip it so the cover page doesn't show a stray rule under
    # the title.
    clear_paragraph_border(document.styles["Title"])

    # Never let a heading strand alone at the bottom of a page with its
    # content pushed to the next one.
    for style_name in ("Heading 1", "Heading 2"):
        document.styles[style_name].paragraph_format.keep_with_next = True


def _build_cover_page(document: Document, *, doc_name: str, document_control: DocumentControlMeta) -> None:
    section = document.sections[0]
    section.different_first_page_header_footer = True

    header_logo = header_logo_path()
    if header_logo.exists():
        header_paragraph = section.first_page_header.paragraphs[0]
        header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_paragraph.add_run().add_picture(str(header_logo), width=Inches(1.3))

    document.add_paragraph()

    project_logo = project_logo_path()
    if project_logo.exists():
        logo_paragraph = document.add_paragraph()
        logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_paragraph.add_run().add_picture(str(project_logo), width=Inches(2.0))
        document.add_paragraph()

    title1 = document.add_paragraph(style="Title")
    title1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title1.add_run("Requirement Analysis Report")

    for_p = document.add_paragraph()
    for_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for_run = for_p.add_run("for")
    for_run.font.name = FONT_NAME

    title2 = document.add_paragraph(style="Title")
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title2.add_run(document_control.title or doc_name)

    if document_control.project_id:
        project_id_p = document.add_paragraph()
        project_id_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        id_run = project_id_p.add_run(f"Project ID: {document_control.project_id}")
        id_run.font.name = FONT_NAME

    # Pushed down to sit near the bottom of the page, not immediately after
    # the title block -- the standard "prepared by" placement on a cover
    # page. This value is measured, not estimated: an analytical estimate
    # from font sizes and image dimensions put the safe ceiling well above
    # where it actually is (Word's default style spacing adds more than
    # expected), and a value chosen from that estimate alone silently
    # overflowed onto an auto-inserted second page ahead of the explicit
    # page_break below -- a real, observed defect caught only by actually
    # rendering it. Bisecting against a real two-line-wrapped 65-character
    # document title (with project logo and Project ID both present, i.e.
    # close to this cover's realistic worst case) via the same Word COM
    # rendering `format-report.sh` uses found the actual overflow threshold
    # at 244pt; 220pt keeps a real margin below that measured value rather
    # than a guessed one, without landing so high it reads as "mid-page"
    # again.
    footer_p = document.add_paragraph()
    footer_p.paragraph_format.space_before = Pt(220)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer_run = footer_p.add_run(f"{document_control.prepared_by}\n{document_control.prepared_date}")
    footer_run.font.name = FONT_NAME

    document.add_page_break()


def _set_content_footer(document: Document) -> None:
    """Page-number footer for every page except the cover (which has none,
    via different_first_page_header_footer)."""
    footer = document.sections[0].footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_p.add_run("Page ")
    run.font.name = FONT_NAME
    run.font.size = Pt(11)
    add_page_field(footer_p, "PAGE", placeholder_text="1")
    run = footer_p.add_run(" of ")
    run.font.name = FONT_NAME
    run.font.size = Pt(11)
    add_page_field(footer_p, "NUMPAGES", placeholder_text="1")
    for footer_run in footer_p.runs:
        footer_run.font.name = FONT_NAME
        footer_run.font.size = Pt(11)


def _add_field(document: Document, label: str, text: str) -> None:
    """One labeled field: a bold label paragraph (glued to the value that
    follows it, so a page break can never strand the label alone with its
    value pushed to the next page) and a plain paragraph holding the value.
    Used for every "Label: value" pair in this report that isn't already a
    controlled-document table field -- the Project Overview & Scope
    narrative and every requirement card's own fields -- so both read
    consistently. The Document Control & Metadata fields use
    `add_key_value_table` instead (see `_build_document_control_section`):
    a controlled-document identification block reads as a table everywhere
    else in this project (the Revision History subsection right below it,
    and the Test Plan's own Document Version Control section), so this
    section follows the same shape rather than standing out as loose
    paragraphs."""
    label_p = document.add_paragraph()
    label_p.paragraph_format.space_before = Pt(10)
    label_p.paragraph_format.space_after = Pt(2)
    label_p.paragraph_format.keep_with_next = True
    label_run = label_p.add_run(label)
    label_run.font.name = FONT_NAME
    label_run.font.size = Pt(FONT_SIZE_PT)
    label_run.font.bold = True
    label_run.font.color.rgb = RGBColor.from_string(ACCENT_COLOR)

    value_p = document.add_paragraph()
    value_p.paragraph_format.space_after = Pt(2)
    value_run = value_p.add_run(text)
    value_run.font.name = FONT_NAME
    value_run.font.size = Pt(FONT_SIZE_PT)


def _build_document_control_section(
    document: Document, *, document_control: DocumentControlMeta, release_history: list[ReleaseHistoryEntry],
    doc_name: str, requirements: list[Requirement],
) -> None:
    document.add_heading("1. Document Control & Metadata", level=1)
    add_key_value_table(
        document,
        [(label, getattr(document_control, attr)) for label, attr in DOCUMENT_CONTROL_FIELDS],
    )

    document.add_heading("1.1 Revision History", level=2)
    revision_rows = [
        [e.version, e.date, e.reasons, e.author, e.reviewed_by, e.approved_by] for e in release_history
    ]
    add_header_table(
        document,
        REVISION_HISTORY_COLUMNS,
        revision_rows,
        REVISION_HISTORY_WIDTHS_IN,
        font_size_pt=REVISION_HISTORY_FONT_SIZE_PT,
        header_fill=HEADER_FILL,
        header_text_color=HEADER_TEXT_COLOR,
        cell_padding_pt=4,
    )

    summary_p = document.add_paragraph()
    summary_run = summary_p.add_run(_build_summary_line(doc_name, requirements))
    summary_run.font.name = FONT_NAME
    summary_run.font.size = Pt(FONT_SIZE_PT)
    summary_run.font.italic = True

    # Front matter never shares a page with the report body.
    document.add_page_break()


def _build_scope_section(document: Document, *, executive_summary: str, in_scope: str, out_of_scope: str) -> None:
    document.add_heading("2. Project Overview & Scope", level=1)
    _add_field(document, "Executive Summary", executive_summary)
    _add_field(document, "In-Scope", in_scope)
    _add_field(document, "Out-of-Scope", out_of_scope)
    document.add_page_break()


def _requirement_heading_text(requirement: Requirement) -> str:
    return f"{requirement.req_id}: {requirement.title}" if requirement.title else requirement.req_id


def _add_requirement_card(document: Document, requirement: Requirement) -> None:
    document.add_heading(_requirement_heading_text(requirement), level=2)

    display_status = _display_status(requirement.acceptance_criteria_status)
    status_p = document.add_paragraph()
    status_p.paragraph_format.space_after = Pt(6)
    status_p.paragraph_format.keep_with_next = True
    status_run = status_p.add_run(f"Status: {display_status}")
    status_run.font.name = FONT_NAME
    status_run.font.size = Pt(FONT_SIZE_PT)
    status_run.font.bold = True
    status_run.font.color.rgb = RGBColor.from_string(STATUS_COLORS.get(display_status, ACCENT_COLOR))

    _add_field(document, "Requirement", requirement.requirement_text)
    _add_field(document, "Gap", _format_gap_text(requirement))
    _add_field(document, "Client Question", _format_question_text(requirement))
    _add_field(document, "Acceptance Criteria", requirement.acceptance_criteria)
    _add_field(document, "Recommendations", _format_recommendations_text(requirement))


def _add_card_separator(document: Document) -> None:
    separator_p = document.add_paragraph()
    separator_p.paragraph_format.space_before = Pt(6)
    separator_p.paragraph_format.space_after = Pt(16)
    add_bottom_border(separator_p)


def _build_category_section(document: Document, category: str, requirements: list[Requirement]) -> None:
    document.add_heading(SECTION_TITLES[category], level=1)
    items = sorted(
        (r for r in requirements if r.category == category),
        key=lambda r: _req_sort_key(r.req_id),
    )
    if not items:
        empty_p = document.add_paragraph()
        empty_run = empty_p.add_run(NO_CATEGORY_REQUIREMENTS_TEXT[category])
        empty_run.font.name = FONT_NAME
        empty_run.font.size = Pt(FONT_SIZE_PT)
        return
    for idx, requirement in enumerate(items):
        _add_requirement_card(document, requirement)
        if idx != len(items) - 1:
            _add_card_separator(document)


def build_report(
    requirements: list[Requirement],
    *,
    doc_name: str,
    document_control: DocumentControlMeta | None = None,
    release_history: list[ReleaseHistoryEntry] | None = None,
    executive_summary: str = "",
    in_scope: str = "",
    out_of_scope: str = "",
) -> Document:
    document_control = document_control or DocumentControlMeta()
    release_history = release_history or []

    document = Document()
    _set_portrait_letter(document)
    # python-docx stamps author="python-docx" into every file it creates; this
    # report is client-facing, where that is readable in File > Info.
    set_authorship(
        document,
        author=document_control.prepared_by or "Emvigo QA",
        title=f"Requirement Analysis — {doc_name}",
        subject=doc_name,
    )

    style = document.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(FONT_SIZE_PT)
    _set_base_styles(document)

    _build_cover_page(document, doc_name=doc_name, document_control=document_control)
    _set_content_footer(document)

    _build_document_control_section(
        document,
        document_control=document_control,
        release_history=release_history,
        doc_name=doc_name,
        requirements=requirements,
    )
    _build_scope_section(
        document, executive_summary=executive_summary, in_scope=in_scope, out_of_scope=out_of_scope
    )

    if not requirements:
        document.add_heading(SECTION_TITLES[CATEGORY_FUNCTIONAL], level=1)
        empty_p = document.add_paragraph()
        empty_run = empty_p.add_run(NO_REQUIREMENTS_TEXT)
        empty_run.font.name = FONT_NAME
        empty_run.font.size = Pt(FONT_SIZE_PT)
    else:
        _build_category_section(document, CATEGORY_FUNCTIONAL, requirements)
        _build_category_section(document, CATEGORY_NON_FUNCTIONAL, requirements)
        _build_category_section(document, CATEGORY_COMPLIANCE, requirements)

    return document


def write_report(doc_name: str) -> None:
    json_path = analysis_json_path(doc_name)
    data = read_json(json_path)

    errors, warnings = validate_file(json_path)
    for warning in warnings:
        print(f"WARNING: {warning}", file=sys.stderr)
    if errors:
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        print(
            f"Refusing to generate the report: {json_path.name} failed validation. "
            "Fix the JSON and re-run.",
            file=sys.stderr,
        )
        sys.exit(1)

    analysis = AnalysisDocument.from_any(data)

    document = build_report(
        analysis.requirements,
        doc_name=doc_name,
        document_control=analysis.document_control,
        release_history=analysis.release_history,
        executive_summary=analysis.meta.executive_summary,
        in_scope=analysis.meta.in_scope,
        out_of_scope=analysis.meta.out_of_scope,
    )

    docx_path = analysis_docx_path(doc_name)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    save_document(document, docx_path)
    print(str(docx_path))


def main() -> None:
    if len(sys.argv) != 2:
        print(
            "Usage: python -m orchestrator.generation.docx_report_writer <doc-name>",
            file=sys.stderr,
        )
        sys.exit(1)

    _, doc_name = sys.argv
    write_report(doc_name)


if __name__ == "__main__":
    main()
