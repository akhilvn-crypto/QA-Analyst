"""Shared exporter: renders a markdown report as a formatted Word document
or PDF.

Usage:
    python -m orchestrator.generation.report_exporter \
        --input <path-to.md> --output-name <basename> --format docx|pdf \
        [--output-dir <dir>]

This is a generic utility, not an agent-specific writer -- any agent that has
already written a markdown report can call it to produce the client-facing
docx/pdf. It reads only the markdown file it is given and writes only the file
it is told to write, so callers stay isolated from each other.

`--output-dir` is optional. When omitted, output goes to `output/reports/`
under the workspace root (the standalone `/export-report` behaviour). When
passed explicitly, the export is written there instead, so a caller can keep
its docx/pdf alongside its own markdown rather than in the generic reports
folder.

Structure is preserved, never flattened: every heading level stays its own
heading, and a labelled field bullet (`- **Finding** — ...`) renders as its own
paragraph with the label bold and intact, so a report's per-finding fields
(Finding / Real Impact / Severity / How to Fix) can never collapse into one
run-together paragraph. Severity values are colour-coded (Critical/High in red
shades, Medium amber, Low green).
"""

import argparse
import re
import shutil
import tempfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Mm, Pt, RGBColor

from orchestrator.utils.docx_helpers import (
    FONT_NAME,
    SEVERITY_STYLES,
    add_page_field,
    bake_live_fields,
    clear_paragraph_border,
    repeat_header_row,
    save_document,
    set_fixed_column_widths,
    scrub_authorship,
    set_authorship,
    set_min_row_height,
    set_row_cant_split,
    set_table_cell_margins,
    set_update_fields_on_open,
    shade_cell,
    style_cell,
)
from orchestrator.utils.paths import exported_reports_dir, project_logo_path

ACCENT_COLOR = "1F3864"  # dark navy -- matches every other report in this project
HEADER_TEXT_COLOR = "FFFFFF"

BODY_SIZE_PT = 11
TABLE_SIZE_PT = 10.5

# A4 portrait with 0.8in side margins => ~6.67in usable. Table widths must sum
# to just under that: Word will not auto-shrink a fixed-layout table that
# overflows the page.
TABLE_TOTAL_WIDTH_IN = 6.6

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
ORDERED_RE = re.compile(r"^(?P<indent>\s*)(?P<number>\d+)[.)]\s+(?P<content>.*)$")
RULE_RE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")

# `- **Label** — value` / `**Label**: value` -- the labelled-field form that
# must survive as a distinct labelled paragraph.
FIELD_RE = re.compile(r"^\*\*(?P<label>[^*]+?)\*\*\s*(?P<sep>[—\-:–])\s*(?P<value>.*)$")

# Inline markdown: bold, italic, inline code.
INLINE_RE = re.compile(r"(\*\*.+?\*\*|__.+?__|\*[^*]+?\*|_[^_]+?_|`[^`]+?`)", re.DOTALL)


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(BODY_SIZE_PT)

    heading_sizes = {"Title": 18, "Heading 1": 14, "Heading 2": 13, "Heading 3": 12, "Heading 4": 11}
    for style_name, size_pt in heading_sizes.items():
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size_pt)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(ACCENT_COLOR)
        # python-docx's built-in Title style ships with a bottom border that
        # isn't a design choice here -- strip it, as the other writers do.
        clear_paragraph_border(style)


def _setup_page(document: Document, *, classification: str | None = None) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)

    # Classification sits on its own line above the page number, so a printed
    # or forwarded page always carries its handling marking.
    if classification:
        marking = section.footer.paragraphs[0]
        marking.alignment = WD_ALIGN_PARAGRAPH.LEFT
        marking_run = marking.add_run(classification)
        marking_run.font.name = FONT_NAME
        marking_run.font.size = Pt(8)
        marking_run.font.italic = True
        marking_run.font.color.rgb = RGBColor.from_string("595959")
        footer = section.footer.add_paragraph()
    else:
        footer = section.footer.paragraphs[0]

    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Page ")
    run.font.name = FONT_NAME
    run.font.size = Pt(9)
    add_page_field(footer, "PAGE", placeholder_text="1")
    run = footer.add_run(" of ")
    run.font.name = FONT_NAME
    run.font.size = Pt(9)
    add_page_field(footer, "NUMPAGES", placeholder_text="1")
    for footer_run in footer.runs:
        footer_run.font.name = FONT_NAME
        footer_run.font.size = Pt(9)
        footer_run.font.color.rgb = RGBColor.from_string("595959")


def _add_cover_page(
    document: Document,
    *,
    title: str,
    subtitle: str | None,
    classification: str | None,
    logo: Path | None,
    prepared_by: str | None,
) -> None:
    """Build the cover page and end it with the document's first page break."""
    # A cover page carries no footer -- classification and page numbering start
    # on the first content page, as in the Test Plan deliverable.
    document.sections[0].different_first_page_header_footer = True

    if logo and logo.is_file():
        picture = document.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(str(logo), width=Inches(2.2))
        document.add_paragraph()

    heading = document.add_paragraph(style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_run = heading.add_run(title)
    heading_run.font.name = FONT_NAME
    heading_run.font.size = Pt(24)
    heading_run.font.bold = True
    heading_run.font.color.rgb = RGBColor.from_string(ACCENT_COLOR)

    if subtitle:
        sub = document.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub.add_run(subtitle)
        sub_run.font.name = FONT_NAME
        sub_run.font.size = Pt(14)
        sub_run.font.color.rgb = RGBColor.from_string("404040")

    if classification:
        document.add_paragraph()
        band = document.add_paragraph()
        band.alignment = WD_ALIGN_PARAGRAPH.CENTER
        band_run = band.add_run(classification.upper())
        band_run.font.name = FONT_NAME
        band_run.font.size = Pt(10)
        band_run.font.bold = True
        band_run.font.color.rgb = RGBColor.from_string("9C0006")

    if prepared_by:
        for _ in range(2):
            document.add_paragraph()
        for label, value in (("Prepared by:", prepared_by), ("Date:", date.today().isoformat())):
            line = document.add_paragraph()
            line.paragraph_format.space_after = Pt(2)
            label_run = line.add_run(f"{label} ")
            label_run.font.name = FONT_NAME
            label_run.font.size = Pt(11)
            label_run.font.bold = True
            value_run = line.add_run(value)
            value_run.font.name = FONT_NAME
            value_run.font.size = Pt(11)

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _add_table_of_contents(document: Document) -> None:
    """Insert a real Word TOC field covering Heading 1-4.

    A live field, not a hand-built list: it stays accurate if a heading's text
    or page changes. `write_report`-style baking (see `bake_live_fields`) fills
    in the computed entries so a client sees them without their Word needing
    update-on-open behaviour.

    The "Table of Contents" caption is styled to match Heading 1 but is
    deliberately *not* a Heading 1 -- a real heading would list the contents
    page as its own first entry.
    """
    heading = document.add_paragraph()
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_after = Pt(12)
    heading_run = heading.add_run("Table of Contents")
    heading_run.font.name = FONT_NAME
    heading_run.font.size = Pt(14)
    heading_run.font.bold = True
    heading_run.font.color.rgb = RGBColor.from_string(ACCENT_COLOR)

    toc = document.add_paragraph()
    add_page_field(
        toc,
        r'TOC \o "1-4" \h \z \u',
        placeholder_text="Right-click and choose “Update Field” to build the contents.",
    )

    # The ToC page has no natural successor heading to break on, so this break
    # is structural, not stylistic -- see the pagination note in
    # .claude/rules/security-audit-output-structure.md.
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def severity_of(text: str) -> str | None:
    """The severity named by `text`, if it names exactly one.

    Matches a bare value ("Critical", "High — 5 findings") but deliberately not
    prose that merely mentions a level in passing, so body text never gets
    colour-coded by accident.
    """
    stripped = re.sub(r"[*_`]", "", text).strip().strip(".:—-").strip()
    if not stripped:
        return None
    first_word = re.split(r"[\s,(/—-]+", stripped)[0].lower()
    if first_word in SEVERITY_STYLES and len(stripped.split()) <= 6:
        return first_word
    return None


def _add_inline_runs(paragraph, text: str, *, size_pt: float = BODY_SIZE_PT, bold: bool = False):
    """Render inline markdown (bold/italic/code) as separate runs."""
    for token in INLINE_RE.split(text):
        if not token:
            continue
        is_bold, is_italic, is_code = bold, False, False
        if token.startswith("**") and token.endswith("**"):
            token, is_bold = token[2:-2], True
        elif token.startswith("__") and token.endswith("__"):
            token, is_bold = token[2:-2], True
        elif token.startswith("*") and token.endswith("*"):
            token, is_italic = token[1:-1], True
        elif token.startswith("_") and token.endswith("_"):
            token, is_italic = token[1:-1], True
        elif token.startswith("`") and token.endswith("`"):
            token, is_code = token[1:-1], True

        run = paragraph.add_run(token)
        run.font.name = "Consolas" if is_code else FONT_NAME
        run.font.size = Pt(size_pt)
        run.font.bold = is_bold
        run.font.italic = is_italic
    return paragraph


def _add_field_paragraph(document: Document, match: re.Match, *, indent_level: int) -> None:
    """Render `**Label** — value` as its own paragraph: bold label, then the
    value. Keeping the label as a distinct bold run is what stops a finding's
    four fields from reading as one collapsed paragraph in the docx/pdf."""
    label = match.group("label").strip()
    value = match.group("value").strip()

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25 + 0.25 * indent_level)
    paragraph.paragraph_format.space_after = Pt(3)

    label_run = paragraph.add_run(f"{label}: ")
    label_run.font.name = FONT_NAME
    label_run.font.size = Pt(BODY_SIZE_PT)
    label_run.font.bold = True
    label_run.font.color.rgb = RGBColor.from_string(ACCENT_COLOR)

    severity = severity_of(value) if label.lower().startswith("severity") else None
    if severity:
        text_color, _ = SEVERITY_STYLES[severity]
        value_run = paragraph.add_run(value)
        value_run.font.name = FONT_NAME
        value_run.font.size = Pt(BODY_SIZE_PT)
        value_run.font.bold = True
        value_run.font.color.rgb = RGBColor.from_string(text_color)
    else:
        _add_inline_runs(paragraph, value)


def _add_numbered_paragraph(
    document: Document, number: str, content: str, *, indent_level: int
) -> None:
    """Render an ordered-list item carrying the markdown's own number.

    Word's built-in `List Number` style auto-numbers continuously for the whole
    document, so a second numbered list would carry on from the first instead
    of restarting at 1. Emitting the source number as literal text keeps every
    list numbered exactly as the markdown states it.
    """
    paragraph = document.add_paragraph()
    left_indent = 0.25 + 0.25 * indent_level
    paragraph.paragraph_format.left_indent = Inches(left_indent)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(3)

    number_run = paragraph.add_run(f"{number}. ")
    number_run.font.name = FONT_NAME
    number_run.font.size = Pt(BODY_SIZE_PT)
    _add_inline_runs(paragraph, content)


def _add_heading(document: Document, level: int, text: str) -> None:
    heading = document.add_heading(level=min(level, 4))
    # Never let a heading be stranded at the bottom of a page with its content
    # pushed to the next one.
    heading.paragraph_format.keep_with_next = True

    severity = severity_of(text)
    color = SEVERITY_STYLES[severity][0] if severity else ACCENT_COLOR
    for run in _add_inline_runs(
        heading, text, size_pt={1: 14, 2: 13, 3: 12, 4: 11}.get(min(level, 4), 11), bold=True
    ).runs:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=column_count)
    table.style = "Table Grid"
    set_table_cell_margins(table)

    for row_index, values in enumerate(rows):
        table_row = table.add_row()
        set_row_cant_split(table_row)
        set_min_row_height(table_row, 16)
        padded = list(values) + [""] * (column_count - len(values))
        for column_index, value in enumerate(padded):
            cell = table_row.cells[column_index]
            clean = re.sub(r"[*_`]", "", value).strip()
            if row_index == 0:
                style_cell(cell, clean, bold=True, color=HEADER_TEXT_COLOR, size_pt=TABLE_SIZE_PT)
                shade_cell(cell, ACCENT_COLOR)
                cell.paragraphs[0].paragraph_format.keep_with_next = True
            else:
                severity = severity_of(clean)
                if severity:
                    text_color, fill = SEVERITY_STYLES[severity]
                    style_cell(cell, clean, bold=True, color=text_color, size_pt=TABLE_SIZE_PT)
                    shade_cell(cell, fill)
                else:
                    style_cell(cell, clean, size_pt=TABLE_SIZE_PT)

        if row_index == 0:
            repeat_header_row(table_row)

    width = round(TABLE_TOTAL_WIDTH_IN / column_count, 3)
    set_fixed_column_widths(table, [width] * column_count)
    document.add_paragraph()


def _is_cover_title(heading_text: str, cover_title: str | None) -> bool:
    """Whether an H1 just restates the cover page's title."""
    if not cover_title:
        return False
    normalise = lambda text: re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()
    return normalise(heading_text) == normalise(cover_title)


def _starts_new_block(line: str) -> bool:
    """Whether `line` begins a new markdown block rather than continuing one."""
    stripped = line.strip()
    if not stripped:
        return True
    return bool(
        HEADING_RE.match(line)
        or BULLET_RE.match(line)
        or ORDERED_RE.match(line)
        or RULE_RE.match(line)
        or stripped.startswith("```")
        or "|" in stripped
    )


def _gather_block(lines: list[str], index: int) -> tuple[str, int]:
    """Join a block's wrapped source lines into one logical line.

    Markdown hard-wraps prose: a single paragraph or bullet is usually several
    source lines that a renderer must join. Treating each source line as its own
    Word paragraph is what makes wrapped prose come out with a blank-ish gap
    after every line, and turns a bullet's continuation into a stray
    unbulleted paragraph. Returns the joined text and the index after the block.
    """
    collected = [lines[index].strip()]
    index += 1
    while index < len(lines) and not _starts_new_block(lines[index]):
        collected.append(lines[index].strip())
        index += 1
    return " ".join(part for part in collected if part), index


def _split_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


@dataclass
class ReportOptions:
    """Optional document furniture that turns a plain export into a
    client-deliverable controlled document.

    Every field defaults to off, so a caller that passes nothing gets the same
    plain output as before -- this is what keeps the standalone
    `/export-report` behaviour unchanged while letting a caller opt into a
    cover page, contents, and classification marking.
    """

    cover_title: str | None = None
    cover_subtitle: str | None = None
    classification: str | None = None
    prepared_by: str | None = None
    logo: Path | None = None
    toc: bool = False

    @property
    def needs_field_baking(self) -> bool:
        """A TOC field carries no cached result until Word computes it."""
        return self.toc

    @property
    def document_author(self) -> str:
        """Who the file's metadata names as author.

        Never the Windows account name and never "python-docx" (python-docx's
        default), both of which a client can read in File > Info.
        """
        return self.prepared_by or "Emvigo QA"


def build_document(markdown: str, options: ReportOptions | None = None) -> Document:
    """Render markdown text as a Word document, preserving its structure."""
    options = options or ReportOptions()

    document = Document()
    _setup_page(document, classification=options.classification)
    _configure_styles(document)
    set_authorship(
        document,
        author=options.document_author,
        title=options.cover_title or "",
        subject=options.cover_subtitle or "",
    )

    if options.cover_title:
        _add_cover_page(
            document,
            title=options.cover_title,
            subtitle=options.cover_subtitle,
            classification=options.classification,
            logo=options.logo,
            prepared_by=options.prepared_by,
        )
    if options.toc:
        _add_table_of_contents(document)

    # With a cover page present, the markdown's own H1 is redundant furniture
    # -- the title is already on the cover -- so let it render as a section
    # heading instead of a second Title block.
    first_heading_used = bool(options.cover_title)

    lines = markdown.replace("\r\n", "\n").split("\n")
    index = 0
    in_code_block = False

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            in_code_block = not in_code_block
            index += 1
            continue

        if in_code_block:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            index += 1
            continue

        if not stripped or RULE_RE.match(line):
            index += 1
            continue

        # Pipe table: a header row followed by a |---|---| separator.
        if (
            "|" in stripped
            and index + 1 < len(lines)
            and TABLE_SEPARATOR_RE.match(lines[index + 1])
        ):
            table_rows = [_split_table_row(stripped)]
            index += 2
            while index < len(lines) and "|" in lines[index].strip():
                table_rows.append(_split_table_row(lines[index]))
                index += 1
            _add_table(document, table_rows)
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and not first_heading_used:
                title = document.add_paragraph(style="Title")
                _add_inline_runs(title, text, size_pt=18, bold=True)
                for run in title.runs:
                    run.font.color.rgb = RGBColor.from_string(ACCENT_COLOR)
                first_heading_used = True
            elif level == 1 and _is_cover_title(text, options.cover_title):
                # The cover page already carries this exact title; repeating it
                # as the first section heading is duplicated furniture.
                pass
            else:
                _add_heading(document, level, text)
            index += 1
            continue

        # Everything left is a bullet, an ordered item, a labelled field, or a
        # paragraph -- each of which may wrap across several source lines, so
        # gather the whole block before deciding how to render it. Indentation
        # comes from the first line, since the gathered text is stripped.
        indent_level = 1 if len(line) - len(line.lstrip()) >= 2 else 0
        block_text, index = _gather_block(lines, index)

        ordered = ORDERED_RE.match(block_text)
        bullet = BULLET_RE.match(block_text)
        if ordered:
            content = ordered.group("content").strip()
        elif bullet:
            content = bullet.group(2).strip()
        else:
            content = block_text

        field = FIELD_RE.match(content)
        if field:
            _add_field_paragraph(document, field, indent_level=indent_level)
        elif ordered:
            _add_numbered_paragraph(
                document, ordered.group("number"), content, indent_level=indent_level
            )
        elif bullet:
            paragraph = document.add_paragraph(
                style="List Bullet 2" if indent_level else "List Bullet"
            )
            _add_inline_runs(paragraph, content)
        else:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(6)
            _add_inline_runs(paragraph, content)

    return document


def _convert_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    """Convert a .docx to PDF via Word COM automation.

    This system only ever runs on Windows machines with Word installed (the
    same assumption `hooks/format-report.sh` makes). Unlike that hook,
    a missing Word is a hard failure here: the PDF is a requested deliverable,
    so silently producing nothing would be worse than saying why.
    """
    try:
        import win32com.client as win32
    except ImportError:
        raise SystemExit(
            "PDF export needs pywin32 and Microsoft Word (Windows only). "
            "Install orchestrator/requirements.txt, or export --format docx instead."
        )

    # Absolutise both paths: Word resolves a relative path against its own
    # working directory, so a valid relative path here fails with "Sorry, we
    # couldn't find your file" (Open) or a bare "Command failed" (SaveAs).
    docx_path = Path(docx_path).resolve()
    pdf_path = Path(pdf_path).resolve()
    pdf_path.parent.mkdir(parents=True, exist_ok=True)

    word = win32.gencache.EnsureDispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(docx_path))
        try:
            # Compute fields before export: a TOC field carries no cached
            # result, so exporting without this produces a PDF whose contents
            # page is blank -- and a PDF can't be updated after the fact.
            doc.Fields.Update()
            for toc in doc.TablesOfContents:
                toc.Update()
            doc.Fields.Update()
            doc.SaveAs(str(pdf_path), FileFormat=17)  # wdFormatPDF
        finally:
            doc.Close(False)
    finally:
        word.Quit()

    if not pdf_path.is_file():
        raise SystemExit(f"Word did not produce {pdf_path}.")


def resolve_output_dir(output_dir: str | None) -> Path:
    """Where the export is written.

    An explicit --output-dir always wins. Otherwise falls back to
    `output/reports/` under the workspace root -- the standalone
    `/export-report` default.
    """
    if output_dir:
        return Path(output_dir)
    return exported_reports_dir()


def export_report(
    input_path: Path,
    output_name: str,
    export_format: str,
    output_dir: str | None = None,
    options: ReportOptions | None = None,
) -> Path:
    if not input_path.is_file():
        raise SystemExit(f"Input markdown not found: {input_path}")

    options = options or ReportOptions()
    destination_dir = resolve_output_dir(output_dir)
    destination_dir.mkdir(parents=True, exist_ok=True)

    markdown = input_path.read_text(encoding="utf-8")
    document = build_document(markdown, options)
    if options.needs_field_baking:
        set_update_fields_on_open(document)

    if export_format == "docx":
        output_path = destination_dir / f"{output_name}.docx"
        save_document(document, output_path)
        if options.needs_field_baking:
            bake_live_fields(output_path)
            # Word's save stamped the Windows account name into
            # last_modified_by; reset it before this file goes to a client.
            scrub_authorship(output_path, author=options.document_author)
        return output_path

    output_path = destination_dir / f"{output_name}.pdf"
    temp_dir = tempfile.mkdtemp(prefix="report-exporter-")
    temp_docx = Path(temp_dir) / f"{output_name}.docx"
    try:
        save_document(document, temp_docx)
        _convert_to_pdf(temp_docx, output_path)
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Export a markdown report as a formatted Word document or PDF."
    )
    parser.add_argument("--input", required=True, help="Path to the source .md file.")
    parser.add_argument(
        "--output-name", required=True, help="Output filename without extension."
    )
    parser.add_argument("--format", required=True, choices=("docx", "pdf"), dest="export_format")
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Destination folder. Defaults to output/reports/ under the workspace root.",
    )

    furniture = parser.add_argument_group(
        "controlled-document options",
        "All optional. Omit them all for a plain export.",
    )
    furniture.add_argument(
        "--cover-title", default=None, help="Render a cover page with this title."
    )
    furniture.add_argument("--cover-subtitle", default=None, help="Cover page subtitle.")
    furniture.add_argument(
        "--classification",
        default=None,
        help='Handling marking, e.g. "Confidential - Client Use Only". Shown on the '
        "cover and in every page footer.",
    )
    furniture.add_argument("--prepared-by", default=None, help="Cover page prepared-by line.")
    furniture.add_argument(
        "--logo",
        default=None,
        help="Cover logo image. Defaults to the project logo asset when --cover-title is set.",
    )
    furniture.add_argument(
        "--toc",
        action="store_true",
        help="Insert a live Word Table of Contents field (Heading 1-4) after the cover.",
    )
    args = parser.parse_args()

    if args.logo:
        logo = Path(args.logo)
    elif args.cover_title:
        logo = project_logo_path()
    else:
        logo = None

    options = ReportOptions(
        cover_title=args.cover_title,
        cover_subtitle=args.cover_subtitle,
        classification=args.classification,
        prepared_by=args.prepared_by,
        logo=logo,
        toc=args.toc,
    )

    output_path = export_report(
        Path(args.input), args.output_name, args.export_format, args.output_dir, options
    )
    print(str(output_path))


if __name__ == "__main__":
    main()
