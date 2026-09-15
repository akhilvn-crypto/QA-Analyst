"""Writes the Client Clarification Sheet — a client-facing table of every
open question from a requirement analysis, with a blank Client Response
column for the client to fill in and send back.

Usage:
    python -m orchestrator.generation.clarification_sheet_writer <doc-name> [--docx] [--force]

Reads output/requirement-analysis/<doc-name>-analysis.json and always writes
output/client-clarifications/<doc-name>-clarifications.md — the round-trip
source `/apply-clarifications` reads back (via
`parsing.clarifications_from_md`). The Word copy
(`<doc-name>-clarifications.docx`) is opt-in via `--docx`, same
always-md/opt-in-docx relationship the requirement-analysis report already
has between `md_report_writer`/`docx_report_writer` — a client or QA can
type an answer directly into either copy; whichever one holds it is what
`/apply-clarifications` ingests.

One row per open question (a requirement with two gaps gets two rows).
Requirements with no gaps don't appear — this sheet is only the open items.

Overwrite guard: if an existing sheet — the `.md`, the `.docx`, or both —
already contains client responses that haven't been ingested yet (rows
whose answer isn't recorded in the analysis JSON's clarification_log), this
writer refuses to regenerate either file unless --force is passed — run
/apply-clarifications first so the answers aren't lost.

The sheet opens with an Obsidian-style YAML frontmatter "Properties" block
(`orchestrator.utils.md_frontmatter`, shared by all four `.md` writers) --
title, document_type, project_id, document_id, version, approved_date,
privacy, tags -- ahead of the `# Client Clarification Sheet` heading, which
is otherwise unchanged. See `build_sheet_md`'s own docstring for where this
sheet's project_id/document_id come from, given it has no document_control
of its own.

Before overwriting an existing `.md` (once the guard above has cleared it),
this writer also archives it to a sibling `history/` folder — same
"preserve every controlled revision" contract `md_report_writer.py` gives
the requirement-analysis report, via the same shared
`orchestrator.utils.md_history` helper. Unlike every other deliverable
here, this sheet carries no version of its own (it has no `meta`/
`document_control`, and is fully regenerated from the analysis each run) —
so its snapshots are always named "unversioned", distinguished from each
other by timestamp alone.
"""

import sys
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt

from orchestrator.models.requirement import AnalysisDocument, DocumentControlMeta
from orchestrator.parsing.clarifications_from_docx import SHEET_COLUMNS
from orchestrator.parsing.clarifications_from_docx import read_sheet as read_sheet_docx
from orchestrator.parsing.clarifications_from_md import read_sheet as read_sheet_md
from orchestrator.utils.docx_helpers import (
    FONT_NAME,
    repeat_header_row,
    save_document,
    set_authorship,
    set_fixed_column_widths,
    set_row_cant_split,
    set_table_cell_margins,
    shade_cell,
    style_cell,
)
from orchestrator.utils.io_helpers import read_json
from orchestrator.utils.md_frontmatter import build_frontmatter, title_with_project
from orchestrator.utils.md_history import snapshot_previous_md
from orchestrator.utils.md_table import build_table
from orchestrator.utils.paths import (
    analysis_json_path,
    clarification_sheet_docx_path,
    clarification_sheet_md_path,
)

# Landscape letter, 10in usable width (matches the analysis report setup).
COLUMN_WIDTHS_IN = [0.9, 3.4, 3.0, 2.7]

FONT_SIZE_PT = 12
HEADER_FILL = "1F3864"
HEADER_TEXT_COLOR = "FFFFFF"

_NOTE_TEXT = (
    "Please type your answer to each question in the Client Response column "
    "and return this document. Each answer unblocks the analysis and testing "
    "of its requirement."
)
_EMPTY_TEXT = (
    "No open clarifications — every question has been answered or no "
    "gaps required client input."
)


def _open_question_rows(analysis: AnalysisDocument) -> list[tuple[str, str, str]]:
    rows = []
    for req in analysis.requirements:
        for gap in req.gaps:
            # A requirement with no genuine gap still carries one Gap entry,
            # holding the "No significant gaps identified." / "None." sentinel
            # (see .claude/rules/output-structure.md) rather than an empty
            # list. That sentinel must never surface as a sheet row -- it
            # isn't a question to send to a client.
            if gap.question.strip().lower() == "none.":
                continue
            rows.append((req.req_id, req.requirement_text, gap.question))
    return rows


def _unprocessed_answers(existing_rows: list[dict], analysis: AnalysisDocument) -> list[str]:
    """Req IDs whose sheet answers aren't in the analysis clarification_log yet."""
    logged = {
        (entry.get("req_id", ""), entry.get("answer", "").strip())
        for entry in analysis.meta.clarification_log
    }
    return [
        row["req_id"]
        for row in existing_rows
        if row["answer"] and (row["req_id"], row["answer"]) not in logged
    ]


def build_sheet(
    rows: list[tuple[str, str, str]], *, doc_name: str
) -> Document:
    document = Document()
    # This sheet is sent directly to the client, so it must not ship
    # python-docx's default author in its metadata.
    set_authorship(
        document,
        author="Emvigo QA",
        title=f"Client Clarification Sheet — {doc_name}",
        subject=doc_name,
    )

    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)

    style = document.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(FONT_SIZE_PT)

    document.add_heading("Client Clarification Sheet", level=1)

    meta = document.add_paragraph()
    meta_run = meta.add_run(
        f"Document: {doc_name}    |    Generated: {date.today().isoformat()}    |    "
        f"Open questions: {len(rows)}"
    )
    meta_run.font.name = FONT_NAME
    meta_run.font.size = Pt(FONT_SIZE_PT)
    meta_run.font.bold = True

    note = document.add_paragraph()
    note_run = note.add_run(_NOTE_TEXT)
    note_run.font.name = FONT_NAME
    note_run.font.size = Pt(FONT_SIZE_PT)
    note_run.font.italic = True

    if not rows:
        empty = document.add_paragraph()
        empty_run = empty.add_run(_EMPTY_TEXT)
        empty_run.font.name = FONT_NAME
        empty_run.font.size = Pt(FONT_SIZE_PT)

    # The table is always rendered (header-only when there are no open
    # questions) so the sheet stays machine-readable for the ingest parser.
    document.add_paragraph()

    table = document.add_table(rows=1, cols=len(SHEET_COLUMNS))
    table.style = "Table Grid"
    set_table_cell_margins(table)

    header_row = table.rows[0]
    repeat_header_row(header_row)
    set_row_cant_split(header_row)
    for idx, column in enumerate(SHEET_COLUMNS):
        cell = header_row.cells[idx]
        style_cell(cell, column, bold=True, color=HEADER_TEXT_COLOR, size_pt=FONT_SIZE_PT)
        shade_cell(cell, HEADER_FILL)
        cell.paragraphs[0].paragraph_format.keep_with_next = True

    for req_id, requirement_text, question in rows:
        table_row = table.add_row()
        set_row_cant_split(table_row)
        cells = table_row.cells
        for idx, value in enumerate((req_id, requirement_text, question, "")):
            style_cell(cells[idx], value, size_pt=FONT_SIZE_PT)

    set_fixed_column_widths(table, COLUMN_WIDTHS_IN)
    return document


def build_sheet_md(
    rows: list[tuple[str, str, str]],
    *,
    doc_name: str,
    document_control: "DocumentControlMeta | None" = None,
) -> str:
    """The `.md` twin of `build_sheet` — same title/meta/note copy, same
    columns, rendered as a Markdown pipe table via `orchestrator.utils.
    md_table` instead of a Word table. Always renders the table (header-only
    when `rows` is empty), same as the docx, so it stays machine-readable
    for `parsing.clarifications_from_md` regardless of how many questions
    are currently open.

    Opens with the same Obsidian-style YAML frontmatter block every other
    `.md` writer here now carries (`orchestrator.utils.md_frontmatter`).
    Unlike the other three, this sheet has no `document_control` of its own
    -- `project_id`/`document_id` are read from the source analysis's
    `document_control` instead (the sheet is derived from that same
    analysis), and `version`/`approved_date` are left blank, matching this
    sheet's own "carries no version of its own" convention (see this
    module's docstring)."""
    document_control = document_control or DocumentControlMeta()
    frontmatter = build_frontmatter(
        title=title_with_project(doc_name, "Client Clarification Sheet"),
        document_type="Client Clarification Sheet",
        tags=["clarification-sheet", "qa"],
        project_id=document_control.project_id,
        document_id=document_control.document_id,
    )

    sections = [
        "# Client Clarification Sheet",
        f"Document: {doc_name}    |    Generated: {date.today().isoformat()}    |    "
        f"Open questions: {len(rows)}",
        _NOTE_TEXT,
    ]
    if not rows:
        sections.append(_EMPTY_TEXT)

    table_rows = [[req_id, requirement_text, question, ""] for req_id, requirement_text, question in rows]
    sections.append(build_table(SHEET_COLUMNS, table_rows))

    return frontmatter + "\n" + "\n\n".join(sections) + "\n"


def write_sheet(doc_name: str, *, docx: bool = False, force: bool = False) -> None:
    json_path = analysis_json_path(doc_name)
    analysis = AnalysisDocument.from_any(read_json(json_path))
    rows = _open_question_rows(analysis)

    md_path = clarification_sheet_md_path(doc_name)
    docx_path = clarification_sheet_docx_path(doc_name)

    if not force:
        pending: set[str] = set()
        held_by = []
        if md_path.exists():
            try:
                found = _unprocessed_answers(read_sheet_md(md_path), analysis)
            except ValueError:
                found = []
            if found:
                pending.update(found)
                held_by.append(md_path.name)
        if docx_path.exists():
            try:
                found = _unprocessed_answers(read_sheet_docx(docx_path), analysis)
            except ValueError:
                found = []
            if found:
                pending.update(found)
                held_by.append(docx_path.name)
        if pending:
            print(
                f"Refusing to overwrite {' and '.join(held_by)}: it contains client "
                f"responses not yet applied to the analysis ({', '.join(sorted(pending))}). "
                "Run /apply-clarifications first, or pass --force to discard them.",
                file=sys.stderr,
            )
            sys.exit(1)

    content = build_sheet_md(rows, doc_name=doc_name, document_control=analysis.document_control)
    md_path.parent.mkdir(parents=True, exist_ok=True)
    snapshot_previous_md(md_path, lambda _content: None)
    md_path.write_text(content, encoding="utf-8")
    print(str(md_path))

    if docx:
        document = build_sheet(rows, doc_name=doc_name)
        docx_path.parent.mkdir(parents=True, exist_ok=True)
        save_document(document, docx_path)
        print(str(docx_path))


def main() -> None:
    flags = {"--docx", "--force"}
    args = [a for a in sys.argv[1:] if a not in flags]
    docx = "--docx" in sys.argv[1:]
    force = "--force" in sys.argv[1:]
    if len(args) != 1:
        print(
            "Usage: python -m orchestrator.generation.clarification_sheet_writer "
            "<doc-name> [--docx] [--force]",
            file=sys.stderr,
        )
        sys.exit(1)

    (doc_name,) = args
    write_sheet(doc_name, docx=docx, force=force)


if __name__ == "__main__":
    main()
